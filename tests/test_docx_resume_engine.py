from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core import docx_resume_engine as engine


def _add_numbering(paragraph, num_id: str = "1") -> None:
    properties = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), num_id)
    numbering.append(level)
    numbering.append(number_id)
    properties.append(numbering)


def _make_resume(path: Path) -> None:
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("SUMMARY")
    document.add_paragraph("Original summary")
    document.add_paragraph("SKILLS")
    heading = document.add_paragraph("Cloud:")
    heading.runs[0].bold = True
    document.add_paragraph("Azure, Python")
    document.add_paragraph("PROFESSIONAL EXPERIENCE")
    document.add_paragraph("Acme Corp, Austin, TX\tJan 2020 - Present")
    document.add_paragraph("Senior Engineer")
    document.add_paragraph("Clients: Example")
    first_bullet = document.add_paragraph("Built Azure systems.")
    _add_numbering(first_bullet)
    second_bullet = document.add_paragraph("Automated Python workflows.")
    _add_numbering(second_bullet)
    document.add_paragraph("CERTIFICATIONS")
    document.add_paragraph("Cert A | Cert B")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("University X")
    document.add_paragraph("B.S. Engineering")
    document.save(path)


def test_parse_limits_edits_to_allowed_sections(tmp_path: Path) -> None:
    resume = tmp_path / "resume.docx"
    _make_resume(resume)

    _, _, _, editable = engine._parse_document(resume)
    kinds = [item["kind"] for item in editable]

    assert kinds == [
        "summary",
        "skill",
        "skill",
        "experience_bullet",
        "experience_bullet",
        "certifications",
    ]
    assert all("Senior Engineer" not in item["original"] for item in editable)
    assert all("University X" not in item["original"] for item in editable)


def test_patch_preserves_locked_paragraphs_and_formatting(tmp_path: Path) -> None:
    resume = tmp_path / "resume.docx"
    output = tmp_path / "output.docx"
    _make_resume(resume)

    _, _, _, editable = engine._parse_document(resume)
    edits = {item["id"]: item["original"] for item in editable}
    edits[editable[0]["id"]] = "Updated summary"

    engine._write_patched_docx(resume, output, edits)
    text = engine.extract_docx_text(output)

    assert "Updated summary" in text
    assert "Senior Engineer" in text
    assert "University X" in text


def test_patch_keeps_non_document_zip_members_byte_identical(tmp_path: Path) -> None:
    resume = tmp_path / "resume.docx"
    output = tmp_path / "output.docx"
    _make_resume(resume)

    _, _, _, editable = engine._parse_document(resume)
    edits = {item["id"]: item["original"] for item in editable}
    engine._write_patched_docx(resume, output, edits)

    with ZipFile(resume) as before, ZipFile(output) as after:
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name)


def test_certifications_can_reorder_but_not_change(tmp_path: Path) -> None:
    resume = tmp_path / "resume.docx"
    output = tmp_path / "output.docx"
    _make_resume(resume)

    _, _, _, editable = engine._parse_document(resume)
    edits = {item["id"]: item["original"] for item in editable}
    certification = next(item for item in editable if item["kind"] == "certifications")

    edits[certification["id"]] = "Cert B | Cert A"
    engine._write_patched_docx(resume, output, edits)
    assert engine._validate_certifications(resume, output) == []

    edits[certification["id"]] = "Cert B | Cert C"
    engine._write_patched_docx(resume, output, edits)
    assert engine._validate_certifications(resume, output)


def test_parse_llm_response_requires_every_edit_id(tmp_path: Path) -> None:
    resume = tmp_path / "resume.docx"
    _make_resume(resume)
    _, _, _, editable = engine._parse_document(resume)
    payload = {
        "edits": [
            {"id": item["id"], "text": item["original"]} for item in editable[:-1]
        ],
        "coverage_terms": [
            {"group": "a", "term": "Engineer", "jd_count": 1, "target_count": 1}
        ],
    }

    with pytest.raises(engine.ResumeGenerationError, match="does not cover every"):
        engine._parse_llm_response(json.dumps(payload), editable)


def test_term_table_counts_exact_strings() -> None:
    terms = [
        {
            "group": "a",
            "term": "Senior Engineer",
            "jd_count": 1,
            "target_count": 1,
        },
        {
            "group": "b",
            "term": "Power Platform",
            "jd_count": 2,
            "target_count": 2,
        },
    ]

    assert engine._validate_term_table(
        "Senior Engineer", "Power Platform and Power Platform", terms
    ) == []

    terms[1]["jd_count"] = 1
    assert engine._validate_term_table(
        "Senior Engineer", "Power Platform and Power Platform", terms
    )


def test_term_table_counts_are_set_deterministically() -> None:
    terms = [
        {"group": "a", "term": "Engineer", "jd_count": 99, "target_count": 1},
        {"group": "b", "term": "Power Platform", "jd_count": 0, "target_count": 1},
        {"group": "e", "term": "Azure DevOps", "jd_count": 7, "target_count": 1},
    ]

    engine._set_deterministic_term_counts(
        "Engineer", "Power Platform and Power Platform", terms
    )

    assert [item["jd_count"] for item in terms] == [1, 2, 0]


def test_coverage_validation_is_literal() -> None:
    terms = [
        {"group": "b", "term": "REST APIs", "jd_count": 1, "target_count": 1}
    ]

    errors, report = engine._validate_coverage("Built REST API integrations", terms)

    assert errors
    assert report[0]["resume_count"] == 0


def test_parse_llm_response_accepts_predicted_skill_group_e(tmp_path: Path) -> None:
    resume = tmp_path / "resume.docx"
    _make_resume(resume)
    _, _, _, editable = engine._parse_document(resume)
    payload = {
        "edits": [{"id": item["id"], "text": item["original"]} for item in editable],
        "coverage_terms": [
            {"group": "a", "term": "Engineer", "jd_count": 1, "target_count": 1},
            {"group": "e", "term": "Solution Architecture", "jd_count": 0, "target_count": 1},
        ],
    }

    parsed = engine._parse_llm_response(json.dumps(payload), editable)

    assert parsed["coverage_terms"][1]["group"] == "e"


def test_predicted_group_e_terms_are_reported_but_not_mandatory() -> None:
    terms = [
        {"group": "e", "term": "Azure DevOps", "jd_count": 0, "target_count": 1}
    ]

    errors, report = engine._validate_coverage("Power Platform", terms)

    assert errors == []
    assert report[0]["resume_count"] == 0


def test_edit_length_validation_rejects_paragraph_expansion() -> None:
    editable = [{"id": "p1", "original": "A" * 100}]

    errors = engine._validate_edit_lengths(editable, {"p1": "B" * 106})

    assert errors == [
        "Paragraph p1 is 106 characters; maximum is 105. "
        "Shorten it by replacing wording"
    ]


def test_edit_length_validation_allows_small_absolute_change_to_short_line() -> None:
    editable = [{"id": "p1", "original": "A" * 15}]

    errors = engine._validate_edit_lengths(editable, {"p1": "B" * 16})

    assert errors == []


def test_invalid_certification_rewrite_is_replaced_with_original(tmp_path: Path) -> None:
    resume = tmp_path / "resume.docx"
    _make_resume(resume)
    _, _, _, editable = engine._parse_document(resume)
    payload = {
        "edits": [
            {
                "id": item["id"],
                "text": "Changed Cert" if item["kind"] == "certifications" else item["original"],
            }
            for item in editable
        ],
        "coverage_terms": [
            {"group": "a", "term": "Engineer", "jd_count": 1, "target_count": 1}
        ],
    }

    parsed = engine._parse_llm_response(json.dumps(payload), editable)
    certification = next(item for item in editable if item["kind"] == "certifications")

    assert parsed["edit_map"][certification["id"]] == certification["original"]


def test_output_name_uses_first_three_job_title_words() -> None:
    assert engine._output_base_name(
        "Senior Power Platform Developer", "Abhishek Keshari"
    ) == "Senior-Power-Platform-Detailed-Resume-Abhishek-Keshari"


def test_windows_libreoffice_is_found_without_environment_variable(
    tmp_path: Path, monkeypatch,
) -> None:
    program_files = tmp_path / "Program Files"
    converter = program_files / "LibreOffice" / "program" / "soffice.exe"
    converter.parent.mkdir(parents=True)
    converter.touch()

    monkeypatch.setattr(engine.os, "name", "nt")
    monkeypatch.setenv("PROGRAMFILES", str(program_files))
    monkeypatch.setenv("PROGRAMFILES(X86)", str(tmp_path / "Program Files (x86)"))
    monkeypatch.delenv("AUTOAPPLY_OFFICE_CONVERTER", raising=False)
    monkeypatch.setattr(engine.shutil, "which", lambda _: None)

    seen = {}

    def fake_run(command, **kwargs):
        seen["converter"] = command[0]
        output = Path(command[command.index("--outdir") + 1]) / "resume.pdf"
        output.write_bytes(b"pdf")
        return type("Result", (), {"returncode": 0, "stderr": "", "stdout": ""})()

    monkeypatch.setattr(engine.subprocess, "run", fake_run)
    source = tmp_path / "resume.docx"
    source.write_bytes(b"docx")

    engine.convert_docx_to_pdf(source, tmp_path / "result.pdf")

    assert seen["converter"] == str(converter)


def test_malformed_llm_json_gets_correction_attempt(
    tmp_path: Path, monkeypatch,
) -> None:
    resume = tmp_path / "resume.docx"
    _make_resume(resume)
    _, _, _, editable = engine._parse_document(resume)
    valid_payload = {
        "edits": [{"id": item["id"], "text": item["original"]} for item in editable],
        "coverage_terms": [
            {"group": "a", "term": "Engineer", "jd_count": 1, "target_count": 1}
        ],
    }
    responses = iter([
        json.dumps({"edits": valid_payload["edits"], "coverage_terms": [{"group": "", "term": ""}]}),
        json.dumps(valid_payload),
    ])
    prompts = []

    def fake_invoke(prompt, _config):
        prompts.append(prompt)
        return next(responses)

    monkeypatch.setattr("core.ai_engine.invoke_llm", fake_invoke)
    monkeypatch.setattr(engine, "convert_docx_to_pdf", lambda _src, dest: dest.write_bytes(b"pdf"))
    monkeypatch.setattr(engine, "_validate_pdf_layout", lambda _master, _generated: [])

    output_docx, output_pdf, metadata = engine.generate_docx_resume(
        master_docx_path=resume,
        job_title="Engineer",
        jd_text="Engineer",
        instructions_text="Keep the resume factual.",
        output_dir=tmp_path / "output",
        applicant_name="Jane Doe",
        llm_config=object(),
    )

    assert output_docx.exists()
    assert output_pdf.exists()
    assert metadata["correction_attempts"] == 1
    assert "Coverage terms must use group a-e" in prompts[1]
